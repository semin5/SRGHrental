"""Generate a one-time MySQL import for the supplied rental-printer workbook."""

from pathlib import Path
import sys
import openpyxl


HEADERS = ("품명", "IP", "MAC", "모델", "규격", "단가", "위치", "위치 (세부)", "비고", "계약일")
COLUMNS = (
    "item_name", "ip_address", "mac_address", "model", "specification",
    "unit_price", "location", "location_detail", "remarks", "contract_date",
)


def sql_value(value):
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        raise ValueError("Unexpected boolean in workbook")
    if isinstance(value, (int, float)):
        return str(value)
    return f"CONVERT(0x{str(value).encode('utf-8').hex()} USING utf8mb4)"


def main(source: Path, destination: Path):
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    if workbook.sheetnames != ["사용", "회수"]:
        raise ValueError(f"Unexpected sheets: {workbook.sheetnames}")
    statements = [
        "CREATE DATABASE srghrental CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;",
        "USE srghrental;",
    ]
    summary = {}
    for sheet_name, table_name, expected_count in (("사용", "in_use", 55), ("회수", "recovered", 6)):
        sheet = workbook[sheet_name]
        rows = list(sheet.values)
        headers = list(rows[1][:10])
        headers[7] = " ".join(headers[7].replace("(", " (").split()) if headers[7] else None
        if tuple(headers) != HEADERS:
            raise ValueError(f"Unexpected header in {sheet_name}: {headers}")
        records = [(row_number, row[:10]) for row_number, row in enumerate(rows, 1)
                   if row_number > 2 and row[0] is not None]
        if len(records) != expected_count:
            raise ValueError(f"{sheet_name}: expected {expected_count} records, found {len(records)}")
        statements.append(f"CREATE TABLE {table_name} ("
                          "source_row INT NOT NULL PRIMARY KEY, "
                          "item_name VARCHAR(255), ip_address VARCHAR(255), mac_address VARCHAR(255), "
                          "model VARCHAR(255), specification VARCHAR(255), unit_price INT, "
                          "location VARCHAR(255), location_detail VARCHAR(255), "
                          "remarks TEXT, contract_date VARCHAR(255)"
                          ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;")
        column_list = "source_row, " + ", ".join(COLUMNS)
        values = ["(" + ", ".join([str(row_number)] + [sql_value(value) for value in record]) + ")"
                  for row_number, record in records]
        statements.append(f"INSERT INTO {table_name} ({column_list}) VALUES\n" + ",\n".join(values) + ";")
        summary[table_name] = (len(records), sum(record[5] or 0 for _, record in records))
    if summary["in_use"] != (55, 2015000):
        raise ValueError(f"Source summary mismatch: {summary['in_use']}")
    destination.write_text("\n".join(statements) + "\n", encoding="utf-8")
    print(summary)


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
